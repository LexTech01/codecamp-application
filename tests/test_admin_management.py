"""Tests for admin user management and bulk applicant actions."""
from io import BytesIO
import pytest
from app import db
from app.models.user import User
from app.models.application import Application
from app.models.notification import Notification
from app.models.cohort import Cohort


def _create_admin(app, email="admin@test.com"):
    with app.app_context():
        u = User(email=email, first_name="Admin", last_name="One", role="admin")
        u.set_password("password123")
        db.session.add(u)
        db.session.commit()
        return u.id


def _create_student(app, email="stu@test.com", first="Stu", last="Dent"):
    with app.app_context():
        u = User(email=email, first_name=first, last_name=last, role="student")
        u.set_password("password123")
        db.session.add(u)
        db.session.commit()
        return u.id


def _login(client, user_id):
    with client.session_transaction() as sess:
        sess["_user_id"] = str(user_id)
        sess["_fresh"] = True
        sess["_sess_v"] = 1


def _submit_application(app, user_id, program="Software Engineering", stage="submitted"):
    with app.app_context():
        a = Application(
            user_id=user_id,
            is_submitted=True,
            pipeline_stage=stage,
            status=stage,
            field_of_study=program,
        )
        db.session.add(a)
        db.session.commit()
        return a.id


# ── User management ─────────────────────────────────────────────


def test_users_page_lists_all(app, client):
    admin_id = _create_admin(app)
    _create_student(app)
    _login(client, admin_id)
    resp = client.get("/admin/users")
    assert resp.status_code == 200
    assert b"stu@test.com" in resp.data


def test_applicant_detail_renders_interview_decision(app, client):
    admin_id = _create_admin(app)
    stu_id = _create_student(app, "det@test.com")
    app_id = _submit_application(app, stu_id, stage="interview_scheduled")
    _login(client, admin_id)
    resp = client.get(f"/admin/applicants/{app_id}")
    assert resp.status_code == 200
    assert b"Mark Passed" in resp.data


def test_interviews_page_renders(client, app):
    admin_id = _create_admin(app, "iv-admin@test.com")
    _login(client, admin_id)
    resp = client.get("/admin/interviews")
    assert resp.status_code == 200
    assert b"Interview Management" in resp.data


def test_promote_student_to_admin(app, client):
    admin_id = _create_admin(app)
    stu_id = _create_student(app)
    _login(client, admin_id)
    resp = client.post(f"/admin/users/{stu_id}/role", data={"role": "admin"})
    assert resp.status_code == 302
    with app.app_context():
        assert User.query.get(stu_id).role == "admin"


def test_demote_student_self_blocked(app, client):
    admin_id = _create_admin(app)
    _login(client, admin_id)
    resp = client.post(f"/admin/users/{admin_id}/role", data={"role": "student"})
    assert resp.status_code == 302
    with app.app_context():
        assert User.query.get(admin_id).role == "admin"


def test_deactivate_student(app, client):
    admin_id = _create_admin(app)
    stu_id = _create_student(app)
    _login(client, admin_id)
    resp = client.post(f"/admin/users/{stu_id}/deactivate", data={"reason": "spam"})
    assert resp.status_code == 302
    with app.app_context():
        u = User.query.get(stu_id)
        assert u.is_active is False
        assert u.deactivated_reason == "spam"
        assert u.deactivated_at is not None


def test_reactivate_student(app, client):
    admin_id = _create_admin(app)
    stu_id = _create_student(app)
    _login(client, admin_id)
    client.post(f"/admin/users/{stu_id}/deactivate", data={"reason": ""})
    resp = client.post(f"/admin/users/{stu_id}/reactivate")
    assert resp.status_code == 302
    with app.app_context():
        u = User.query.get(stu_id)
        assert u.is_active is True
        assert u.deactivated_reason is None


def test_deactivated_user_cannot_be_admin_acted_on_self(app, client):
    admin_id = _create_admin(app)
    _login(client, admin_id)
    resp = client.post(f"/admin/users/{admin_id}/deactivate", data={"reason": ""})
    assert resp.status_code == 302
    with app.app_context():
        assert User.query.get(admin_id).is_active is True


# ── Bulk actions ────────────────────────────────────────────────


def test_bulk_move_to_under_review(app, client):
    admin_id = _create_admin(app)
    s1 = _create_student(app, "s1@test.com")
    s2 = _create_student(app, "s2@test.com")
    a1 = _submit_application(app, s1)
    a2 = _submit_application(app, s2)
    _login(client, admin_id)
    resp = client.post("/admin/applicants/bulk-action", data={
        "app_ids": [str(a1), str(a2)],
        "action": "move_to_under_review",
    })
    assert resp.status_code == 302
    with app.app_context():
        assert Application.query.get(a1).pipeline_stage == "under_review"
        assert Application.query.get(a2).pipeline_stage == "under_review"


def test_bulk_reject_with_reason(app, client):
    admin_id = _create_admin(app)
    s1 = _create_student(app, "r1@test.com")
    a1 = _submit_application(app, s1)
    _login(client, admin_id)
    resp = client.post("/admin/applicants/bulk-action", data={
        "app_ids": [str(a1)],
        "action": "reject",
        "reason": "Not eligible",
    })
    assert resp.status_code == 302
    with app.app_context():
        app_obj = Application.query.get(a1)
        assert app_obj.pipeline_stage == "rejected"
        assert app_obj.rejection_reason == "Not eligible"


def test_bulk_action_creates_notification(app, client):
    admin_id = _create_admin(app)
    s1 = _create_student(app, "n@test.com")
    a1 = _submit_application(app, s1)
    _login(client, admin_id)
    client.post("/admin/applicants/bulk-action", data={
        "app_ids": [str(a1)],
        "action": "move_to_under_review",
    })
    with app.app_context():
        assert Notification.query.filter_by(user_id=s1).count() == 1


def test_bulk_action_no_selection(app, client):
    admin_id = _create_admin(app)
    _login(client, admin_id)
    resp = client.post("/admin/applicants/bulk-action", data={"action": "reject"})
    assert resp.status_code == 302


def test_bulk_action_bad_ids_no_500(app, client):
    admin_id = _create_admin(app)
    _login(client, admin_id)
    resp = client.post("/admin/applicants/bulk-action", data={
        "app_ids": ["not-a-number", "999999"],
        "action": "move_to_under_review",
    })
    assert resp.status_code == 302


# ── Interview pass/fail/waitlist decisions ─────────────────────────────


def test_reinstate_rejected_to_test(app, client):
    admin_id = _create_admin(app)
    stu_id = _create_student(app, "reins@test.com")
    app_id = _submit_application(app, stu_id, stage="rejected")
    _login(client, admin_id)
    resp = client.post(f"/admin/applicants/{app_id}/update-stage",
                       data={"pipeline_stage": "test_invited"})
    assert resp.status_code == 302
    with app.app_context():
        assert Application.query.get(app_id).pipeline_stage == "test_invited"


def test_interview_mark_passed_moves_to_accepted(app, client):
    admin_id = _create_admin(app)
    stu_id = _create_student(app, "pass@test.com")
    app_id = _submit_application(app, stu_id, stage="interview_scheduled")
    _login(client, admin_id)
    resp = client.post(f"/admin/applicants/{app_id}/update-stage",
                       data={"pipeline_stage": "accepted"})
    assert resp.status_code == 302
    with app.app_context():
        assert Application.query.get(app_id).pipeline_stage == "accepted"


def test_interview_mark_failed_rejects(app, client):
    admin_id = _create_admin(app)
    stu_id = _create_student(app, "fail@test.com")
    app_id = _submit_application(app, stu_id, stage="interview_scheduled")
    _login(client, admin_id)
    resp = client.post(f"/admin/applicants/{app_id}/update-stage", data={
        "pipeline_stage": "rejected",
        "rejection_reason": "Did not meet interview bar",
    })
    assert resp.status_code == 302
    with app.app_context():
        app_obj = Application.query.get(app_id)
        assert app_obj.pipeline_stage == "rejected"
        assert app_obj.rejection_reason == "Did not meet interview bar"


def test_interview_cannot_waitlist(app, client):
    admin_id = _create_admin(app)
    stu_id = _create_student(app, "wl@test.com")
    app_id = _submit_application(app, stu_id, stage="interview_scheduled")
    _login(client, admin_id)
    resp = client.post(f"/admin/applicants/{app_id}/update-stage",
                       data={"pipeline_stage": "waitlisted"})
    assert resp.status_code == 302
    with app.app_context():
        assert Application.query.get(app_id).pipeline_stage == "interview_scheduled"


# ── Cohort assignment without JS ───────────────────────────────────────


def test_assign_unassigned_form_driven(app, client):
    admin_id = _create_admin(app)
    stu_id = _create_student(app, "asg@test.com")
    app_id = _submit_application(app, stu_id, stage="accepted")
    with app.app_context():
        cohort = Cohort(name="Cohort Alpha")
        db.session.add(cohort)
        db.session.commit()
        cohort_id = cohort.id
    _login(client, admin_id)
    resp = client.post("/admin/cohorts/assign-unassigned", data={
        "cohort_id": str(cohort_id),
        "app_ids": [str(app_id)],
    })
    assert resp.status_code == 302
    with app.app_context():
        app_obj = Application.query.get(app_id)
        assert app_obj.cohort_name == "Cohort Alpha"
        assert app_obj.pipeline_stage == "onboarding"


def test_assign_unassigned_bad_cohort(app, client):
    admin_id = _create_admin(app)
    _login(client, admin_id)
    resp = client.post("/admin/cohorts/assign-unassigned", data={
        "cohort_id": "0",
        "app_ids": ["1"],
    })
    assert resp.status_code == 404


# ── Spreadsheet export ──────────────────────────────────────────

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _read_export(resp):
    from openpyxl import load_workbook
    return load_workbook(BytesIO(resp.data)).active


def test_export_applicants_xlsx(app, client):
    admin_id = _create_admin(app)
    s1 = _create_student(app, "csv@test.com", first="Csv", last="User")
    a1 = _submit_application(app, s1, program="Networking")
    _login(client, admin_id)
    resp = client.get("/admin/applicants/export.xlsx")
    assert resp.status_code == 200
    assert XLSX_MIME in resp.headers["Content-Type"]
    ws = _read_export(resp)
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    assert rows[0][0] == "First Name"
    assert any(r[2] == "csv@test.com" for r in rows)
    # Columns auto-expand so full headers are fully visible
    from openpyxl.utils import get_column_letter
    for i, name in enumerate(rows[0], start=1):
        width = ws.column_dimensions[get_column_letter(i)].width
        assert width is not None and width >= len(str(name))


def test_export_applicants_xlsx_contains_application_details(app, client):
    admin_id = _create_admin(app)
    s1 = _create_student(app, "full@test.com", first="Full", last="Details")
    with app.app_context():
        a = Application(
            user_id=s1,
            is_submitted=True,
            pipeline_stage="test_invited",
            status="test_invited",
            date_of_birth="2000-01-01",
            age=26,
            gender="Female",
            city="Lagos",
            country="Nigeria",
            referral_source="Friend",
            campus_location="CodeCamp Accra",
            field_of_study="Software Engineering",
            referral_code="FAYA",
            previous_student="No, I'm new to Cellusys CodeCamp",
            current_status="Graduate",
            profession="Computer Science",
            institution="Unilag",
            why_cellusys="Great program",
            career_goals="Become an engineer",
            can_commit=True,
        )
        db.session.add(a)
        db.session.commit()
    _login(client, admin_id)
    resp = client.get("/admin/applicants/export.xlsx")
    ws = _read_export(resp)
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    header = [str(v) for v in rows[0]]
    assert "Date of Birth" in header
    assert "Referral Source" in header
    assert "Campus" in header
    assert "Can Commit" in header
    assert "Why Cellusys" in header
    assert "Career Goals" in header
    assert len(header) == 22
    for inactive in (
        "Motivation", "Languages Known", "Address",
        "Nationality", "Education Level", "GPA", "Stage", "Test Score",
        "Submitted At", "Cellusys Student", "Applying For",
    ):
        assert inactive not in header
    row = next(r for r in rows if r[2] == "full@test.com")
    data = dict(zip(header, [str(v) if v is not None else "" for v in row]))
    assert data["First Name"] == "Full"
    assert data["WhatsApp Number"] == ""
    assert data["Date of Birth"] == "2000-01-01"
    assert data["Gender"] == "Female"
    assert data["Referral Source"] == "Friend"
    assert data["Campus"] == "CodeCamp Accra"
    assert data["Course"] == "Software Engineering"
    assert data["Referral Code"] == "FAYA"
    assert data["Current Status"] == "Graduate"
    assert data["Profession"] == "Computer Science"
    assert data["Institution"] == "Unilag"
    assert data["Why Cellusys"] == "Great program"
    assert data["Career Goals"] == "Become an engineer"
    assert data["Can Commit"] == "Yes"
    # Widths sized to fit the longest header + value
    assert ws.column_dimensions["B"].width >= 8  # Last Name
    assert ws.column_dimensions["C"].width >= 13  # full@test.com email


def test_export_respects_stage_filter(app, client):
    admin_id = _create_admin(app)
    s1 = _create_student(app, "keep@test.com")
    s2 = _create_student(app, "out@test.com")
    a1 = _submit_application(app, s1, stage="submitted")
    _submit_application(app, s2, stage="rejected")
    _login(client, admin_id)
    resp = client.get("/admin/applicants/export.xlsx?stage=submitted")
    ws = _read_export(resp)
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    emails = [str(r[2]) for r in rows[1:]]
    assert "keep@test.com" in emails
    assert "out@test.com" not in emails


def test_export_cohort_xlsx_handles_special_name(app, client):
    admin_id = _create_admin(app)
    s1 = _create_student(app, "roster@test.com", first="Ro", last="Ster")
    with app.app_context():
        db.session.add(Application(
            user_id=s1, is_submitted=True, pipeline_stage="accepted",
            status="accepted", cohort_name="Cohort: [Special]/One",
        ))
        db.session.add(Cohort(
            name="Cohort: [Special]/One",
            description="Tests export with Excel-forbidden characters",
        ))
        db.session.commit()
    _login(client, admin_id)
    resp = client.get("/admin/cohorts/1/export")
    assert resp.status_code == 200
    assert XLSX_MIME in resp.headers["Content-Type"]
    assert "Cohort__[Special]_One_roster.xlsx" in resp.headers["Content-Disposition"]
    ws = _read_export(resp)
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    assert rows[0][0] == "First Name"
    assert any(r[2] == "roster@test.com" for r in rows)
