"""Spreadsheet export helpers shared by admin applicant/cohort exports.

Contains only the applicant details actively collected by the application
form (see app/routes/student.py ``_save_application_draft``) plus Why
Cellusys and Career Goals. System and administrative attributes (pipeline
stage, test scores, interviews, cohort, timestamps) are intentionally
excluded.

Exports are generated as real ``.xlsx`` workbooks so columns can expand and
every header shows in full, which plain CSV cannot guarantee. If openpyxl is
missing from the runtime environment, exports fall back to CSV so the button
never breaks.
"""
from io import BytesIO

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:  # pragma: no cover - environment-specific
    OPENPYXL_AVAILABLE = False
    Workbook = Font = get_column_letter = None

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MIME = "text/csv; charset=utf-8"


def _yes_no(value):
    return "Yes" if value else "No"


def sanitize_sheet_title(title):
    """Make a tab title Excel-safe (no ``[]:*?/\\``, max 31 chars)."""
    safe = title or "Sheet"
    for ch in "[]:*?/\\":
        safe = safe.replace(ch, "-")
    safe = safe.strip()[:31] or "Sheet"
    return safe


def application_columns():
    return [
        # Step 1 — Personal Details
        "First Name", "Last Name", "Email", "WhatsApp Number",
        "Date of Birth", "Age", "Gender", "City", "Country",
        "Referral Source",
        # Step 2 — Programme Selection
        "Campus", "Course", "Referral Code",
        # Step 3 — Educational / Professional Background
        "Previous Student", "Previous Course", "Previous Cohort",
        "Current Status", "Profession", "Institution", "Why Cellusys",
        "Career Goals", "Can Commit",
    ]


def application_row(app_record):
    a = app_record
    user = a.user
    return [
        # Step 1 — Personal Details
        user.first_name,
        user.last_name,
        user.email,
        user.phone or "",
        a.date_of_birth or "",
        a.age if a.age is not None else "",
        a.gender or "",
        a.city or "",
        a.country or "",
        a.referral_source or "",
        # Step 2 — Programme Selection
        a.campus_location or "",
        a.field_of_study or "",
        a.referral_code or "",
        # Step 3 — Educational / Professional Background
        a.previous_student or "",
        a.previous_course or "",
        a.previous_cohort or "",
        a.current_status or "",
        a.profession or "",
        a.institution or "",
        a.why_cellusys or "",
        a.career_goals or "",
        _yes_no(a.can_commit) if a.can_commit is not None else "",
    ]


def build_applications_xlsx(apps, sheet_title="Applicants"):
    """Build an Excel workbook with auto-expanding column widths.

    Column widths expand to fit the header and the longest value in each
    column so every header is fully visible. The header row is bold and
    frozen for easy reading.
    """
    columns = application_columns()
    rows = [columns] + [application_row(a) for a in apps]

    wb = Workbook()
    ws = wb.active
    ws.title = sanitize_sheet_title(sheet_title)

    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)

    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"

    for c in range(1, len(columns) + 1):
        longest = 0
        for row in rows:
            value = row[c - 1]
            if value is not None:
                longest = max(longest, len(str(value)))
        ws.column_dimensions[get_column_letter(c)].width = min(longest + 3, 70)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_applications_csv(apps):
    """Fallback plain-CSV export (used when openpyxl is unavailable)."""
    import csv
    from io import StringIO

    output = StringIO()
    writer = csv.writer(output, lineterminator="\r\n")
    writer.writerow(application_columns())
    for a in apps:
        writer.writerow(application_row(a))
    return ("\ufeff" + output.getvalue()).encode("utf-8")


def build_export(apps, sheet_title="Applicants"):
    """Return ``(bytes, extension, mimetype)`` for the download response.

    Prefers the auto-fitted Excel workbook; falls back to CSV only when
    openpyxl is missing from the runtime.
    """
    if OPENPYXL_AVAILABLE:
        return build_applications_xlsx(apps, sheet_title), "xlsx", XLSX_MIME
    return build_applications_csv(apps), "csv", CSV_MIME